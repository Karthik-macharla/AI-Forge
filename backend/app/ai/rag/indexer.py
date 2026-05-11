"""RAG document indexer.

Splits a file into chunks, embeds them via the LiteLLM proxy, and stores
them in a persistent ChromaDB collection keyed to the attachment ID.

Supported filetypes: PDF (text-based or image/scanned), plain text, Python,
JavaScript, Excel. Images and videos are skipped.
For image-based PDFs, pages are rendered with pymupdf and then OCR'd via
the Gemini vision model through the LiteLLM proxy.
"""
import asyncio
import base64
import uuid
from pathlib import Path

from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_chroma import Chroma
from langchain_core.documents import Document

from app.ai.llm import embeddings, openai_client
from app.core.config import settings
from app.core.logging import logger

# ── Chunk parameters ───────────────────────────────────────────────────────────
_CHUNK_SIZE = 1000
_CHUNK_OVERLAP = 150

# ── ChromaDB collection name (single shared collection per deployment) ─────────
COLLECTION_NAME = "amzur_rag"


def _get_vectorstore() -> Chroma:
    """Return the persistent ChromaDB vectorstore (creates if absent)."""
    return Chroma(
        collection_name=COLLECTION_NAME,
        embedding_function=embeddings,
        persist_directory=settings.CHROMA_PERSIST_DIR,
    )


def _pdf_extract_text_pypdf(path: Path) -> str:
    """Try to extract text from a PDF using pypdf. Returns empty string if none."""
    try:
        import pypdf
        reader = pypdf.PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        logger.warning("pypdf extraction failed for %s: %s", path, exc)
        return ""


def _pdf_render_pages_as_b64(path: Path, dpi: int = 150) -> list[str]:
    """Render each PDF page to a PNG image and return as base64 strings.
    Uses pymupdf (fitz), which handles both text and image-based PDFs.
    """
    import fitz  # pymupdf
    doc = fitz.open(str(path))
    pages_b64 = []
    mat = fitz.Matrix(dpi / 72, dpi / 72)  # scale factor vs default 72dpi
    for page in doc:
        pix = page.get_pixmap(matrix=mat)
        png_bytes = pix.tobytes("png")
        pages_b64.append(base64.b64encode(png_bytes).decode("utf-8"))
    doc.close()
    return pages_b64


def _ocr_page_via_vision(b64_png: str) -> str:
    """Send a single page image to the vision model and get back extracted text."""
    response = openai_client.chat.completions.create(
        model=settings.LLM_MODEL,
        messages=[
            {
                "role": "user",
                "content": [
                    {
                        "type": "text",
                        "text": (
                            "Extract ALL text from this document page exactly as it appears, "
                            "preserving table structure with tab-separated columns where relevant. "
                            "Output only the extracted text with no commentary."
                        ),
                    },
                    {
                        "type": "image_url",
                        "image_url": {"url": f"data:image/png;base64,{b64_png}"},
                    },
                ],
            }
        ],
        max_tokens=4096,
    )
    return response.choices[0].message.content or ""


async def _ocr_pdf_via_vision(path: Path) -> str:
    """Render all pages and OCR them via the vision model. Returns combined text."""
    logger.info("Image-based PDF detected — using vision OCR for %s", path.name)
    pages_b64 = await asyncio.to_thread(_pdf_render_pages_as_b64, path)
    page_texts = []
    for i, b64 in enumerate(pages_b64):
        logger.info("OCR page %d/%d of %s", i + 1, len(pages_b64), path.name)
        text = await asyncio.to_thread(_ocr_page_via_vision, b64)
        page_texts.append(text)
    return "\n\n".join(page_texts)


def _load_text(file_path: str, mime_type: str) -> str:
    """Extract raw text from a file based on its MIME type.
    For PDFs this only handles text-based PDFs — image PDFs are handled
    separately via _ocr_pdf_via_vision (async).
    """
    path = Path(file_path)

    if mime_type == "application/pdf":
        return _pdf_extract_text_pypdf(path)

    if mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(path), data_only=True)
            lines = []
            for sheet in wb.worksheets:
                lines.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    lines.append("\t".join(str(c) if c is not None else "" for c in row))
            return "\n".join(lines)
        except Exception as exc:
            logger.warning("Excel parse failed for %s: %s", file_path, exc)
            return ""

    # text/plain, text/x-python, text/javascript, etc.
    return path.read_text(encoding="utf-8", errors="replace")


async def index_attachment(
    attachment_id: str,
    file_path: str,
    mime_type: str,
    file_name: str,
    user_id: str,
    thread_id: str,
) -> int:
    """Chunk, embed, and store a document in ChromaDB.

    Returns the number of chunks indexed (0 if skipped).
    """
    # Skip binary types that the LLM handles via vision / video endpoints
    if mime_type.startswith("image/") or mime_type.startswith("video/"):
        logger.info("Skipping RAG indexing for %s (type=%s)", file_name, mime_type)
        return 0

    logger.info("RAG indexing: %s (%s)", file_name, mime_type)

    # Load text — for PDFs try pypdf first, fall back to vision OCR if empty
    raw_text = await asyncio.to_thread(_load_text, file_path, mime_type)

    if not raw_text.strip() and mime_type == "application/pdf":
        # Image-based / scanned PDF — use vision model for OCR
        try:
            raw_text = await _ocr_pdf_via_vision(Path(file_path))
        except Exception as exc:
            logger.error("Vision OCR failed for %s: %s", file_name, exc)

    if not raw_text.strip():
        logger.warning("No text extracted from %s — skipping RAG index", file_name)
        return 0

    # Split into overlapping chunks
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=_CHUNK_SIZE,
        chunk_overlap=_CHUNK_OVERLAP,
        separators=["\n\n", "\n", " ", ""],
    )
    chunks = splitter.split_text(raw_text)
    if not chunks:
        return 0

    # Build LangChain Documents with rich metadata for later filtering
    docs = [
        Document(
            page_content=chunk,
            metadata={
                "attachment_id": attachment_id,
                "file_name": file_name,
                "mime_type": mime_type,
                "user_id": user_id,
                "thread_id": thread_id,
                "chunk_index": i,
            },
        )
        for i, chunk in enumerate(chunks)
    ]

    # Generate deterministic IDs so re-indexing the same attachment is idempotent
    doc_ids = [f"{attachment_id}_{i}" for i in range(len(docs))]

    # Embed + store (runs sync Chroma in thread pool to avoid blocking async loop)
    def _store():
        vs = _get_vectorstore()
        vs.add_documents(docs, ids=doc_ids)

    await asyncio.to_thread(_store)
    logger.info("RAG indexed %d chunks for attachment %s", len(chunks), attachment_id)
    return len(chunks)


async def delete_attachment_index(attachment_id: str) -> None:
    """Remove all ChromaDB chunks for a given attachment."""
    def _delete():
        vs = _get_vectorstore()
        vs.delete(where={"attachment_id": attachment_id})

    await asyncio.to_thread(_delete)
    logger.info("RAG index deleted for attachment %s", attachment_id)
